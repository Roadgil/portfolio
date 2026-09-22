import os
import re
import math
import copy
import subprocess
import shutil
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from datetime import datetime
from pathlib import Path

import pandas as pd
from bs4 import BeautifulSoup

import ci_invoice_excel


# ── 기본 경로 ────────────────────────────────────────────
DEFAULT_TEMPLATE_HTML  = r"C:\Users\yoongil.chae\OneDrive - Candela\Syneron-Candela Korea - Operation\12. 수출\rebalance_invoice_logo.html"
DEFAULT_TEMPLATE_EXCEL = ci_invoice_excel.DEFAULT_EXCEL_TEMPLATE
DEFAULT_IMPORT        = r"C:\Users\yoongil.chae\OneDrive - Candela\Syneron-Candela Korea - Operation\10. 수입\수입신고실적(20250916~).xlsx"
DEFAULT_OUTPUT        = r"C:\Users\yoongil.chae\OneDrive - Candela\Syneron-Candela Korea - Operation\12. 수출\CI"
DEFAULT_EXPORT_EXCEL  = r"C:\Users\yoongil.chae\OneDrive - Candela\Syneron-Candela Korea - Operation\12. 수출\수출신고실적20260126~.xlsx"
IMPORT_SHEET          = "수입실적(202509~)"
EXPORT_SHEET          = "Sheet1"

RMA_PRICE_RATIO = 0.10
RMA_HTS_CODE = "9801.00.1090"   # RMA(반송품) HTS code 고정값
PDF_SCALE = 0.78

EXPORT_COL = {
    "destination": 1,
    "purpose":     2,
    "export_date": 3,
    "bl":          4,
    "hs_code":     5,
    "part_code":   6,
    "description": 7,
    "qty":         8,
    "unit_price":  9,
    "amount":      10,
    "so":          11,
    "delivery":    12,
    "origin":      13,
}


# ── 유틸 ─────────────────────────────────────────────────
def normalize_text(v):
    if v is None:
        return ""
    if isinstance(v, float) and math.isnan(v):
        return ""
    s = str(v).strip()
    return "" if s.lower() == "nan" else s


def generate_search_key(v):
    s = normalize_text(v)
    if s.endswith(".0"):
        s = s[:-2]
    return re.sub(r"[^A-Z0-9]", "", s.upper())


def safe_float(v):
    if v is None:
        return None
    if isinstance(v, (int, float)) and not pd.isna(v):
        return float(v)
    s = normalize_text(v).replace(",", "")
    try:
        return float(s) if s else None
    except Exception:
        return None


def format_qty(v):
    if v is None:
        return ""
    f = float(v)
    if abs(f - int(f)) < 1e-9:
        return str(int(f))
    return f"{f:g}"


def format_money(v, decimals=2):
    if v is None:
        return ""
    return f"{float(v):,.{decimals}f}"


def normalize_sheet_name(name):
    return re.sub(r"\s+", " ", str(name).replace("\xa0", " ").strip()).upper()


def compact_desc(part, desc, max_len=None):
    part = normalize_text(part)
    desc = normalize_text(desc)
    if not desc:
        return part
    text = f"{part} {desc}"
    if max_len is None:
        return text
    if len(text) <= max_len:
        return text
    return text[:max_len].rstrip() + "..."


def find_browser_executable():
    candidates = [
        shutil.which("msedge"),
        shutil.which("chrome"),
        shutil.which("google-chrome"),
        shutil.which("chromium"),
        r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
        r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
        r"C:\Program Files\Google\Chrome\Application\chrome.exe",
        r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
    ]
    for path in candidates:
        if path and os.path.exists(path):
            return path
    return None


def render_html_to_pdf(html_path, pdf_path):
    try:
        from playwright.sync_api import sync_playwright

        with sync_playwright() as p:
            browser = p.chromium.launch()
            page = browser.new_page(viewport={"width": 1800, "height": 1100})
            page.goto(Path(html_path).resolve().as_uri(), wait_until="networkidle")
            page.pdf(
                path=pdf_path,
                format="A4",
                landscape=True,
                print_background=True,
                margin={"top": "8mm", "right": "8mm", "bottom": "8mm", "left": "8mm"},
                prefer_css_page_size=True,
            )
            browser.close()
        return
    except Exception:
        pass

    browser = find_browser_executable()
    if not browser:
        raise Exception(
            "PDF 렌더러를 찾지 못했습니다.\n"
            "1) playwright 설치: pip install playwright && playwright install chromium\n"
            "또는\n"
            "2) Microsoft Edge / Google Chrome 설치"
        )

    cmd = [
        browser,
        "--headless=new",
        "--disable-gpu",
        "--run-all-compositor-stages-before-draw",
        f"--print-to-pdf={pdf_path}",
        "--print-to-pdf-no-header",
        Path(html_path).resolve().as_uri(),
    ]
    result = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", errors="replace")
    if result.returncode != 0 or not os.path.exists(pdf_path):
        raise Exception(
            "브라우저 PDF 변환 실패\n"
            f"CMD: {' '.join(cmd)}\n"
            f"STDERR: {result.stderr.strip()}"
        )


def write_to_export_excel(excel_path, rows):
    import openpyxl
    from openpyxl.styles import Alignment

    wb = openpyxl.load_workbook(excel_path)
    if EXPORT_SHEET not in wb.sheetnames:
        raise ValueError(f"'{EXPORT_SHEET}' 시트가 없습니다.")
    ws = wb[EXPORT_SHEET]

    last_row = 1
    for row in ws.iter_rows(min_row=2, max_col=max(EXPORT_COL.values())):
        if any(c.value is not None for c in row):
            last_row = row[0].row
    next_row = last_row + 1

    center = Alignment(horizontal="center", vertical="center")

    for r in rows:
        ws.cell(next_row, EXPORT_COL["destination"]).value = r["destination"]
        ws.cell(next_row, EXPORT_COL["purpose"]).value     = r["purpose"]
        ws.cell(next_row, EXPORT_COL["export_date"]).value = r["export_date"]
        ws.cell(next_row, EXPORT_COL["bl"]).value          = "N/A"
        ws.cell(next_row, EXPORT_COL["hs_code"]).value     = r["hs_code"] or ""
        ws.cell(next_row, EXPORT_COL["part_code"]).value   = r["part_code"]
        ws.cell(next_row, EXPORT_COL["description"]).value = r["description"]
        ws.cell(next_row, EXPORT_COL["qty"]).value         = r["qty"]
        ws.cell(next_row, EXPORT_COL["unit_price"]).value  = r["unit_price"]
        ws.cell(next_row, EXPORT_COL["amount"]).value      = f"=I{next_row}*H{next_row}"
        ws.cell(next_row, EXPORT_COL["so"]).value          = r["so"]
        ws.cell(next_row, EXPORT_COL["delivery"]).value    = r["delivery"]
        ws.cell(next_row, EXPORT_COL["origin"]).value      = r["origin"]

        for col_idx in EXPORT_COL.values():
            ws.cell(next_row, col_idx).alignment = center

        next_row += 1

    wb.save(excel_path)
    return next_row - (last_row + 1)


# ── HTML 처리 ────────────────────────────────────────────
def load_html_soup(path):
    html = Path(path).read_text(encoding="utf-8", errors="ignore")
    return BeautifulSoup(html, "html.parser")


def extract_sheet_map_from_html(path):
    html = Path(path).read_text(encoding="utf-8", errors="ignore")
    pattern = re.compile(
        r'<button[^>]*onclick="showSheet\((\d+)\)"[^>]*>(.*?)</button>',
        re.IGNORECASE | re.DOTALL
    )

    mapping = {}
    ordered = []

    for m in pattern.finditer(html):
        idx = int(m.group(1))
        raw_name = re.sub(r"<.*?>", "", m.group(2))
        raw_name = raw_name.replace("\xa0", " ").strip()
        norm_name = normalize_sheet_name(raw_name)
        mapping[norm_name] = idx
        ordered.append((idx, raw_name))

    return mapping, ordered


def get_html_sheet_names(html_path):
    _, ordered = extract_sheet_map_from_html(html_path)
    return [name for _, name in ordered]


def find_sheet_index_by_name(soup, sheet_name, html_path=None):
    target = normalize_sheet_name(sheet_name)

    if html_path:
        mapping, _ = extract_sheet_map_from_html(html_path)
        if target in mapping:
            return mapping[target]

    tabs = soup.select(".tabs .tab")
    for i, btn in enumerate(tabs):
        if normalize_sheet_name(btn.get_text(" ", strip=True)) == target:
            return i

    available = []
    if html_path:
        _, ordered = extract_sheet_map_from_html(html_path)
        available = [name for _, name in ordered]

    raise Exception(
        f"HTML 템플릿에서 시트 '{sheet_name}' 을(를) 찾지 못했습니다.\n"
        f"사용 가능 시트: {available}"
    )


def get_sheet_by_name(soup, sheet_name, html_path=None):
    idx = find_sheet_index_by_name(soup, sheet_name, html_path=html_path)

    sheet = soup.find("div", id=f"sheet-{idx}")
    if sheet is not None:
        return sheet, idx

    sheets = soup.select(".sheet")
    if idx < len(sheets):
        return sheets[idx], idx

    raise Exception(f"시트 인덱스 {idx} 에 해당하는 .sheet 가 없습니다.")


def get_table_rows(sheet_tag):
    return sheet_tag.select("tr")


def cell_text(td):
    return td.get_text(" ", strip=True).replace("\xa0", " ").strip()


def row_text(tr):
    return tr.get_text(" ", strip=True).replace("\xa0", " ").strip()


def tr_cells_with_positions(tr):
    pos = 1
    out = []
    for td in tr.find_all("td", recursive=False):
        colspan = int(td.get("colspan", "1"))
        out.append((td, pos, colspan))
        pos += colspan
    return out


def cell_at_column(tr, target_col):
    for td, start, colspan in tr_cells_with_positions(tr):
        if start == target_col:
            return td
    return None


def find_row_index_containing(rows, text):
    target = normalize_sheet_name(text)
    for i, tr in enumerate(rows):
        if target in normalize_sheet_name(row_text(tr)):
            return i
    return -1


def find_cell_by_exact_text(rows, text):
    target = normalize_sheet_name(text)
    for tr in rows:
        for td, start, colspan in tr_cells_with_positions(tr):
            if normalize_sheet_name(cell_text(td)) == target:
                return tr, td, start, colspan
    return None, None, None, None


def set_cell_value(td, value):
    td.clear()
    td.string = "" if value is None else str(value)


def ensure_pdf_css(soup):
    css = f"""
@page {{
  size: A4 landscape;
  margin: 8mm;
}}

html, body {{
  background: white !important;
  margin: 0 !important;
  padding: 0 !important;
  width: 100%;
  height: 100%;
  overflow: visible !important;
  font-family: Arial, sans-serif !important;
}}

.tabs {{
  display: none !important;
}}

.sheet {{
  display: none !important;
  padding: 0 !important;
  margin: 0 !important;
}}

.sheet.active {{
  display: flex !important;
  justify-content: center !important;
  padding: 0 !important;
  margin: 0 !important;
}}

.sheet-wrap {{
  box-shadow: none !important;
  background: white !important;
  padding: 0 !important;
  margin: 0 !important;
  width: 1360px !important;
  zoom: {PDF_SCALE} !important;
}}

.sheet-table {{
  border-collapse: collapse !important;
  table-layout: fixed !important;
  width: 1360px !important;
}}

.sheet-table td {{
  white-space: normal !important;
  word-break: break-word !important;
  overflow: visible !important;
  vertical-align: top !important;
  padding: 2px 3px !important;
  font-size: 12px !important;
  line-height: 1.25 !important;
  height: auto !important;
}}

.sheet-table tr {{
  height: auto !important;
}}

img {{
  max-height: 52px !important;
  width: auto !important;
}}
"""
    style = soup.new_tag("style")
    style.string = css
    if soup.head:
        soup.head.append(style)
    else:
        soup.insert(0, style)


def activate_only_selected_sheet(soup, sheet_name, html_path=None):
    idx = find_sheet_index_by_name(soup, sheet_name, html_path=html_path)

    for tab in soup.select(".tabs .tab"):
        classes = [c for c in tab.get("class", []) if c != "active"]
        onclick = tab.get("onclick", "")
        m = re.search(r"showSheet\((\d+)\)", onclick)
        if m and int(m.group(1)) == idx:
            classes.append("active")
        tab["class"] = classes

    for sheet in soup.select(".sheet"):
        if sheet.get("id") == f"sheet-{idx}":
            classes = [c for c in sheet.get("class", []) if c != "active"]
            classes.append("active")
            sheet["class"] = classes
        else:
            sheet.decompose()

    tabs = soup.select_one(".tabs")
    if tabs:
        tabs.decompose()


def fill_header_fields(sheet, delivery_no, today_str):
    rows = get_table_rows(sheet)

    # NUMBER / DELIVERY
    label_row, _, start_col, _ = find_cell_by_exact_text(rows, "NUMBER / DELIVERY")
    if label_row is not None:
        row_idx = rows.index(label_row)
        if row_idx + 1 < len(rows):
            target_td = cell_at_column(rows[row_idx + 1], start_col)
            if target_td:
                set_cell_value(target_td, delivery_no)

    # DATE
    label_row, _, start_col, _ = find_cell_by_exact_text(rows, "DATE")
    if label_row is not None:
        row_idx = rows.index(label_row)
        if row_idx + 1 < len(rows):
            target_td = cell_at_column(rows[row_idx + 1], start_col)
            if target_td:
                set_cell_value(target_td, today_str)

    # SHIP DATE는 DATE와 동일 값으로 기입
    label_row, _, start_col, _ = find_cell_by_exact_text(rows, "SHIP DATE")
    if label_row is not None:
        row_idx = rows.index(label_row)
        if row_idx + 1 < len(rows):
            target_td = cell_at_column(rows[row_idx + 1], start_col)
            if target_td:
                set_cell_value(target_td, today_str)


def locate_item_area(sheet):
    rows = get_table_rows(sheet)
    header_idx = find_row_index_containing(rows, "SO NO.")
    export_idx = find_row_index_containing(rows, "EXPORT DECLARATION:")

    if header_idx < 0 or export_idx < 0:
        raise Exception("템플릿에서 품목 영역 또는 EXPORT DECLARATION 영역을 찾지 못했습니다.")

    start_idx = header_idx + 2
    if start_idx + 1 >= export_idx:
        raise Exception("템플릿의 품목 블록 구조를 찾지 못했습니다.")

    template_row_1 = copy.deepcopy(rows[start_idx])

    return rows, header_idx, export_idx, template_row_1


def remove_item_row_borders(tr):
    """
    품목 본문 행만 세로 구분선을 제거한다.
    - 헤더는 건드리지 않음
    - 본문은 내부 세로선 제거
    - 맨 왼쪽(SO NO.) / 맨 오른쪽(EXTENSION) 외곽선은 유지
    """
    tds = tr.find_all("td", recursive=False)

    style = tr.get("style", "")
    if style and not style.strip().endswith(";"):
        style += ";"
    style += "border-top:none !important; border-bottom:none !important;"
    tr["style"] = style

    for idx, td in enumerate(tds):
        td_style = td.get("style", "")
        if td_style and not td_style.strip().endswith(";"):
            td_style += ";"

        # 본문 가로선 제거
        td_style += "border-top:none !important; border-bottom:none !important;"

        # 내부 세로선 제거, 좌/우 외곽선만 유지
        if idx == 0:
            td_style += "border-left:1px solid #000 !important; border-right:none !important;"
        elif idx == len(tds) - 1:
            td_style += "border-left:none !important; border-right:1px solid #000 !important;"
        else:
            td_style += "border-left:none !important; border-right:none !important;"

        td["style"] = td_style


def rebuild_item_rows(sheet, items, so_no):
    rows, header_idx, export_idx, row_tpl1 = locate_item_area(sheet)

    for tr in rows[header_idx + 2:export_idx]:
        tr.decompose()

    rows = get_table_rows(sheet)
    export_idx = find_row_index_containing(rows, "EXPORT DECLARATION:")
    export_tr = rows[export_idx]

    for idx, item in enumerate(items, start=1):
        tr1 = copy.deepcopy(row_tpl1)
        tds1 = tr1.find_all("td", recursive=False)

        for td in tds1:
            set_cell_value(td, "")

        # description은 자르지 않고 전체 표시
        # 길면 셀 안에서 자동 줄바꿈되고, 행 높이가 함께 늘어나도록 둔다
        part_desc = compact_desc(item["part"], item["desc"], max_len=None)

        if len(tds1) >= 11:
            set_cell_value(tds1[0], so_no)
            set_cell_value(tds1[1], idx)
            set_cell_value(tds1[2], "")
            set_cell_value(tds1[3], part_desc)

            set_cell_value(tds1[4], item["origin"])
            set_cell_value(tds1[5], item["hts"])
            set_cell_value(tds1[6], format_qty(item["qty"]))
            set_cell_value(tds1[7], "0")
            set_cell_value(tds1[8], format_qty(item["qty"]))
            set_cell_value(tds1[9], format_money(item["price"], 4))
            set_cell_value(tds1[10], format_money(item["amount"], 2))

            # Country of Origin: 현재 맞는 상태라 유지
            style = tds1[4].get("style", "")
            if style and not style.strip().endswith(";"):
                style += ";"
            style += (
                "text-align:center !important;"
                "padding-left:10px !important;"
                "padding-right:0px !important;"
            )
            tds1[4]["style"] = style

            # HTS Code ~ Extension: left 정렬 + 약한 padding-left로 같은 블록을 왼쪽 정렬
            for idx_col in [5, 6, 7, 8, 9, 10]:
                style = tds1[idx_col].get("style", "")
                if style and not style.strip().endswith(";"):
                    style += ";"
                style += (
                    "text-align:left !important;"
                    "padding-left:6px !important;"
                    "padding-right:0px !important;"
                )
                tds1[idx_col]["style"] = style

        remove_item_row_borders(tr1)
        export_tr.insert_before(tr1)


def fill_totals(sheet, total_amount):
    rows = get_table_rows(sheet)
    export_idx = find_row_index_containing(rows, "EXPORT DECLARATION:")
    if export_idx < 0 or export_idx + 1 >= len(rows):
        return

    label_row = rows[export_idx]
    value_row = rows[export_idx + 1]

    # SUBTOTAL / TOTAL 레이블의 실제 컬럼 위치를 찾아서 값 행의 대응 셀에 기입
    # (시트마다 colspan 구조가 달라도 올바른 위치를 찾을 수 있음)
    subtotal_col = None
    total_col = None
    for td, start, colspan in tr_cells_with_positions(label_row):
        text = normalize_sheet_name(cell_text(td))
        if text == "SUBTOTAL":
            subtotal_col = start
        elif text == "TOTAL":
            total_col = start

    for td, start, colspan in tr_cells_with_positions(value_row):
        if subtotal_col is not None and start == subtotal_col:
            set_cell_value(td, format_money(total_amount, 2))
        if total_col is not None and start == total_col:
            set_cell_value(td, format_money(total_amount, 2))


def build_pdf_html(template_html_path, sheet_name, delivery_no, so_no, items, output_html_path):
    soup = load_html_soup(template_html_path)
    activate_only_selected_sheet(soup, sheet_name, html_path=template_html_path)
    ensure_pdf_css(soup)

    sheet, _ = get_sheet_by_name(soup, sheet_name, html_path=template_html_path)

    today_str = datetime.now().strftime("%Y-%m-%d")
    fill_header_fields(sheet, delivery_no, today_str)
    rebuild_item_rows(sheet, items, so_no)

    total_amount = sum(float(x["amount"]) for x in items)
    fill_totals(sheet, total_amount)

    Path(output_html_path).write_text(str(soup), encoding="utf-8")
    return total_amount


# ── 앱 ──────────────────────────────────────────────────
class CIApp:
    def __init__(self, root):
        self.root = root
        self.root.title("CI Generator (HTML PDF Final)")
        self.root.geometry("1200x960")

        self.template_path       = tk.StringVar(value=DEFAULT_TEMPLATE_HTML)
        self.template_excel_path = tk.StringVar(value=DEFAULT_TEMPLATE_EXCEL)
        self.import_path         = tk.StringVar(value=DEFAULT_IMPORT)
        self.output_dir          = tk.StringVar(value=DEFAULT_OUTPUT)
        self.export_excel_path   = tk.StringVar(value=DEFAULT_EXPORT_EXCEL)

        self.country  = tk.StringVar()
        self.delivery = tk.StringVar()
        self.so       = tk.StringVar()
        self.mode     = tk.StringVar(value="Rebalance")
        self.output_format = tk.StringVar(value="EXCEL")

        self.lookup = {}
        self.template_sheets = []

        self.build_ui()
        self.tree.bind("<Double-1>", self.on_double_click)
        self.mode.trace_add("write", lambda *a: self._refresh_mode_ui())
        self._refresh_mode_ui()

    def build_ui(self):
        mf = ttk.LabelFrame(self.root, text=" 📌 문서 유형 선택 ", padding=10)
        mf.pack(fill="x", padx=10, pady=(8, 0))
        ttk.Radiobutton(mf, text="Rebalance  (일반 수출 CI)",
                        variable=self.mode, value="Rebalance").pack(side="left", padx=20)
        ttk.Radiobutton(mf, text="RMA  (시트 고정 FA LAB / 단가 10%)",
                        variable=self.mode, value="RMA").pack(side="left", padx=20)
        self.mode_label = ttk.Label(mf, text="", foreground="#27ae60",
                                    font=("맑은 고딕", 9, "bold"))
        self.mode_label.pack(side="left", padx=20)

        of = ttk.LabelFrame(self.root, text=" 🧾 출력 서식 선택 ", padding=10)
        of.pack(fill="x", padx=10, pady=(8, 0))
        ttk.Radiobutton(of, text="HTML → PDF (기존 방식)",
                        variable=self.output_format, value="HTML").pack(side="left", padx=20)
        ttk.Radiobutton(of, text="Excel 원본 서식 (Rebalance Invoice form)",
                        variable=self.output_format, value="EXCEL").pack(side="left", padx=20)

        top = ttk.LabelFrame(self.root, text=" ⚙️ 기본 설정 ", padding=10)
        top.pack(fill="x", padx=10, pady=5)
        top.columnconfigure(1, weight=1)

        def _file_row(r, label, var, pick_cmd):
            ttk.Label(top, text=label).grid(row=r, column=0, sticky="w")
            ttk.Entry(top, textvariable=var, width=100).grid(
                row=r, column=1, padx=5, pady=2, sticky="ew")
            ttk.Button(top, text="찾기", command=pick_cmd).grid(row=r, column=2)

        _file_row(0, "CI HTML 템플릿:", self.template_path,
                  lambda: self.template_path.set(
                      filedialog.askopenfilename(filetypes=[("HTML", "*.html;*.htm")])))

        _file_row(1, "CI Excel 서식:", self.template_excel_path,
                  lambda: self.template_excel_path.set(
                      filedialog.askopenfilename(filetypes=[("Excel", "*.xlsx")])))

        _file_row(2, "수입 실적:", self.import_path,
                  lambda: self.import_path.set(
                      filedialog.askopenfilename(filetypes=[("Excel", "*.xlsx")])))

        _file_row(3, "수출신고실적:", self.export_excel_path,
                  lambda: self.export_excel_path.set(
                      filedialog.askopenfilename(filetypes=[("Excel", "*.xlsx")])))

        ttk.Button(top, text="🔄 1단계: 수입 데이터 로드 (Cache)",
                   command=self.load_cache).grid(row=4, column=1, sticky="w", pady=8)

        mid = ttk.Frame(self.root, padding=10)
        mid.pack(fill="x")

        inf = ttk.Frame(mid)
        inf.pack(side="left", fill="y", padx=(0, 20))

        ttk.Label(inf, text="국가/시트명:").grid(row=0, column=0, sticky="w")
        self.country_entry = ttk.Entry(inf, textvariable=self.country, width=20)
        self.country_entry.grid(row=0, column=1, pady=5)
        self.auto_match_btn = ttk.Button(inf, text="자동매칭",
                                         command=self.auto_match_country)
        self.auto_match_btn.grid(row=0, column=2, padx=5)

        ttk.Label(inf, text="Delivery No:").grid(row=1, column=0, sticky="w")
        ttk.Entry(inf, textvariable=self.delivery, width=20).grid(row=1, column=1, pady=5)

        ttk.Label(inf, text="SO No:").grid(row=2, column=0, sticky="w")
        ttk.Entry(inf, textvariable=self.so, width=20).grid(row=2, column=1, pady=5)

        itf = ttk.LabelFrame(mid, text=" 품목 입력 (자재코드 수량, 한 줄씩) ", padding=5)
        itf.pack(side="left", fill="both", expand=True)
        self.items_text = tk.Text(itf, width=50, height=8, font=("Consolas", 10))
        self.items_text.pack(fill="both", expand=True)

        btn_frame = ttk.Frame(self.root, padding=(10, 0))
        btn_frame.pack(fill="x")

        ttk.Button(btn_frame, text="🔍 2단계: 미리보기",
                   command=self.preview_items).pack(side="left", padx=5)
        ttk.Button(btn_frame, text="📄 3단계: PDF 생성",
                   command=self.create_ci_pdf).pack(side="left", padx=5)
        ttk.Button(btn_frame, text="📊 4단계: 수출신고실적 기입",
                   command=self.write_export_excel).pack(side="left", padx=5)

        ttk.Separator(btn_frame, orient="vertical").pack(side="left", padx=10, fill="y")
        ttk.Label(btn_frame, text="출력 폴더:").pack(side="left")
        ttk.Entry(btn_frame, textvariable=self.output_dir, width=50).pack(side="left", padx=5)
        ttk.Button(btn_frame, text="찾기",
                   command=lambda: self.output_dir.set(filedialog.askdirectory())).pack(side="left")

        cols = ("자재코드", "품명", "원산지", "HTS", "수량", "단가", "금액", "상태")
        self.tree = ttk.Treeview(self.root, columns=cols, show="headings", height=12)
        for c, w in zip(cols, [130, 260, 60, 100, 60, 90, 90, 160]):
            self.tree.heading(c, text=c)
            self.tree.column(c, width=w, anchor="center")
        self.tree.pack(fill="both", expand=True, padx=10, pady=5)

        log_frame = ttk.LabelFrame(self.root, text=" 로그 ", padding=5)
        log_frame.pack(fill="x", padx=10, pady=(0, 8))
        self.log_box = tk.Text(log_frame, height=7, state="normal", font=("Consolas", 9))
        self.log_box.pack(fill="x")

    def _refresh_mode_ui(self):
        if self.mode.get() == "RMA":
            self.country.set("FA LAB")
            self.country_entry.config(state="disabled")
            self.auto_match_btn.config(state="disabled")
            self.mode_label.config(
                text="※ RMA 모드: 시트 고정(FA LAB) / 단가 10% 적용",
                foreground="#c0392b")
        else:
            self.country_entry.config(state="normal")
            self.auto_match_btn.config(state="normal")
            self.mode_label.config(
                text="※ Rebalance 모드: 일반 로직 적용",
                foreground="#27ae60")

    def _effective_price(self, base_price):
        if base_price is None:
            base_price = 0.0
        return base_price * RMA_PRICE_RATIO if self.mode.get() == "RMA" else base_price

    def on_double_click(self, event):
        item_id = self.tree.identify_row(event.y)
        column  = self.tree.identify_column(event.x)
        if not item_id or not column:
            return

        col_idx = int(column.replace("#", "")) - 1
        if col_idx not in [0, 1, 2, 3, 4, 5]:
            return

        bbox = self.tree.bbox(item_id, column)
        if not bbox:
            return

        x, y, w, h = bbox
        values = list(self.tree.item(item_id, "values"))
        var = tk.StringVar(value=values[col_idx])

        entry = tk.Entry(self.tree, textvariable=var, font=("Consolas", 10))
        entry.place(x=x, y=y, width=w, height=h)
        entry.focus_set()
        entry.select_range(0, "end")

        def commit(e=None):
            values[col_idx] = var.get().strip()
            try:
                q = float(values[4])
                p = float(str(values[5]).replace(",", ""))
                values[6] = f"{q * p:.2f}"
                values[7] = "수정됨"
            except Exception:
                pass
            self.tree.item(item_id, values=values)
            entry.destroy()
            self.update_total_log()

        entry.bind("<Return>", commit)
        entry.bind("<Tab>", commit)
        entry.bind("<Escape>", lambda e: entry.destroy())
        entry.bind("<FocusOut>", commit)

    def update_total_log(self):
        total = sum(safe_float(self.tree.item(r, "values")[6]) or 0
                    for r in self.tree.get_children())
        self.log(f"💰 [{self.mode.get()}] 현재 합계: ₩{total:,.0f} KRW")

    def log(self, msg):
        self.log_box.insert("end", msg + "\n")
        self.log_box.see("end")

    def load_cache(self):
        try:
            df = pd.read_excel(self.import_path.get(), sheet_name=IMPORT_SHEET, header=0)
            # 실적 파일 시트의 "사용된 범위"가 실제 데이터보다 훨씬 아래(100만 행 이상)까지
            # 부풀어 있어(엑셀 서식 잔재), 자재코드가 없는 빈 행을 먼저 걸러내지 않으면
            # 1단계 로드가 20초 넘게 걸린다. 로직은 동일(빈 행은 원래도 걸러지던 행)하고
            # 속도만 개선하는 것.
            df = df[df["자재코드"].notna()]
            hs_col = next(
                (c for c in df.columns
                 if str(c).lower().replace(" ", "").replace("_", "") == "hscode"),
                None
            )
            loaded = 0
            for _, row in df.iloc[::-1].iterrows():
                sk = generate_search_key(row.get("자재코드"))
                if sk and sk not in self.lookup:
                    hs_raw = row.get(hs_col) if hs_col else None
                    if hs_raw is not None and not (isinstance(hs_raw, float) and math.isnan(hs_raw)):
                        hs_str = str(hs_raw).strip().split(".")[0]
                    else:
                        hs_str = ""
                    self.lookup[sk] = {
                        "desc":   normalize_text(row.get("자재명")),
                        "origin": normalize_text(row.get("원산지코드")),
                        "hts":    hs_str,
                        "price":  safe_float(row.get("단가")),
                    }
                    loaded += 1

            messagebox.showinfo("완료", f"데이터 로드 완료 ({loaded}건)\nHS code 컬럼: '{hs_col}'")
        except Exception as e:
            messagebox.showerror("Error", str(e))

    def auto_match_country(self):
        try:
            if self.output_format.get() == "EXCEL":
                import openpyxl
                wb = openpyxl.load_workbook(self.template_excel_path.get(), read_only=True)
                self.template_sheets = [s for s in wb.sheetnames if s != "Address"]
            else:
                self.template_sheets = get_html_sheet_names(self.template_path.get())
        except Exception as e:
            messagebox.showerror("Error", f"템플릿 시트 읽기 실패\n{e}")
            return

        raw = self.country.get().strip().upper()
        alias = {"AU": "AUP", "JP": "JPP", "CN": "CHP", "IL": "ILH", "US": "US"}
        target = alias.get(raw, raw)

        for s in self.template_sheets:
            if normalize_sheet_name(s) == normalize_sheet_name(target):
                self.country.set(s)
                return

    def preview_items(self):
        if not self.lookup:
            messagebox.showwarning("주의", "먼저 1단계를 실행하세요.")
            return

        for r in self.tree.get_children():
            self.tree.delete(r)

        items, _ = self.parse_items()
        is_rma = self.mode.get() == "RMA"

        for item in items:
            info       = self.lookup.get(item["search_key"], {})
            desc       = info.get("desc", "N/A")
            origin     = info.get("origin", "")
            hts        = info.get("hts", "")
            base_price = info.get("price") or 0.0
            price      = self._effective_price(base_price)
            ext        = item["qty"] * price
            status     = "OK" if info else "조회불가"

            # RMA(반송품)는 lookup 결과와 무관하게 HTS code 고정
            if is_rma:
                hts = RMA_HTS_CODE

            if is_rma and info:
                status = f"RMA(10%: {base_price:.2f}→{price:.2f})"

            self.tree.insert("", "end", values=(
                item["raw_part"], desc, origin, hts,
                item["qty"], f"{price:.4f}", f"{ext:.2f}", status
            ))

        self.update_total_log()

    def parse_items(self):
        items = []
        for line in self.items_text.get("1.0", "end").strip().splitlines():
            line = line.strip()
            if not line:
                continue
            if re.search(r"[\t,]", line):
                parts = re.split(r"[\t,]", line, maxsplit=1)
            else:
                tokens = line.rsplit(None, 1)
                parts = tokens if len(tokens) == 2 else []
            if len(parts) >= 2:
                raw = parts[0].strip()
                sk  = generate_search_key(raw)
                qty = safe_float(parts[1].strip())
                if sk and qty is not None:
                    items.append({"raw_part": raw, "search_key": sk, "qty": qty})
        return items, []

    def _selected_sheet_name(self):
        return "FA LAB" if self.mode.get() == "RMA" else self.country.get().strip()

    def _collect_pdf_items(self):
        row_ids = self.tree.get_children()
        items = []
        for row_id in row_ids:
            vals = self.tree.item(row_id, "values")
            qty = safe_float(vals[4]) or 0.0
            price = safe_float(vals[5]) or 0.0
            amount = safe_float(vals[6]) or (qty * price)
            items.append({
                "part": vals[0],
                "desc": vals[1],
                "origin": vals[2],
                "hts": vals[3],
                "qty": qty,
                "price": price,
                "amount": amount,
            })
        return items

    def create_ci_pdf(self):
        sheet_name  = self._selected_sheet_name()
        delivery_no = self.delivery.get().strip()
        so_no       = self.so.get().strip()

        if not sheet_name or not delivery_no:
            messagebox.showerror("Error", "입력 정보를 확인하세요.")
            return

        row_ids = self.tree.get_children()
        if not row_ids:
            messagebox.showerror("Error", "미리보기 데이터가 없습니다.")
            return

        if self.output_format.get() == "EXCEL":
            self._create_ci_pdf_excel(sheet_name, delivery_no, so_no)
        else:
            self._create_ci_pdf_html(sheet_name, delivery_no, so_no)

    def _create_ci_pdf_excel(self, sheet_name, delivery_no, so_no):
        mode_tag  = f"({self.mode.get()})"
        base_name = f"KRP - {sheet_name} - {delivery_no} {mode_tag}".replace("/", "-")
        work_dir  = os.path.abspath(self.output_dir.get())
        xlsx_final = os.path.join(work_dir, f"{base_name}.xlsx")
        pdf_final  = os.path.join(work_dir, f"{base_name}.pdf")

        try:
            os.makedirs(work_dir, exist_ok=True)
            items = self._collect_pdf_items()

            ci_invoice_excel.build_ci_excel(
                template_path=self.template_excel_path.get(),
                sheet_name=sheet_name,
                delivery_no=delivery_no,
                so_no=so_no,
                items=items,
                output_xlsx_path=xlsx_final,
            )
            ci_invoice_excel.export_sheet_to_pdf(xlsx_final, pdf_final, sheet_name=sheet_name)

            self.log(f"[DEBUG] Excel/PDF 생성 완료: {xlsx_final}")
            messagebox.showinfo(
                "성공",
                f"Excel 서식 기반 CI 생성 완료! ({self.mode.get()} 모드)\n"
                f"Excel: {xlsx_final}\nPDF: {pdf_final}"
            )
        except Exception as e:
            self.log(f"[ERROR] {e}")
            messagebox.showerror("Error", str(e))

    def _create_ci_pdf_html(self, sheet_name, delivery_no, so_no):
        mode_tag  = f"({self.mode.get()})"
        base_name = f"KRP - {sheet_name} - {delivery_no} {mode_tag}".replace("/", "-")
        work_dir  = os.path.abspath(self.output_dir.get())
        pdf_final = os.path.join(work_dir, f"{base_name}.pdf")
        html_temp = os.path.join(work_dir, f"{base_name}.html")

        try:
            os.makedirs(work_dir, exist_ok=True)

            items = self._collect_pdf_items()
            total_amount = build_pdf_html(
                template_html_path=self.template_path.get(),
                sheet_name=sheet_name,
                delivery_no=delivery_no,
                so_no=so_no,
                items=items,
                output_html_path=html_temp
            )

            if os.path.exists(pdf_final):
                try:
                    os.remove(pdf_final)
                except Exception:
                    pass

            render_html_to_pdf(html_temp, pdf_final)

            if not os.path.exists(pdf_final):
                raise Exception(f"PDF 생성 실패: {pdf_final}")

            self.log("[DEBUG] PDF 생성 완료")
            messagebox.showinfo(
                "성공",
                f"PDF 생성 완료! ({self.mode.get()} 모드)\n경로: {pdf_final}"
            )

        except Exception as e:
            self.log(f"[ERROR] {e}")
            messagebox.showerror("Error", str(e))

        finally:
            try:
                if os.path.exists(html_temp):
                    os.remove(html_temp)
                    self.log(f"[DEBUG] 임시 HTML 삭제: {html_temp}")
            except Exception as delete_err:
                self.log(f"[WARN] 임시 HTML 삭제 실패: {delete_err}")

    def write_export_excel(self):
        row_ids = self.tree.get_children()
        if not row_ids:
            messagebox.showwarning("주의",
                "미리보기 데이터가 없습니다. 2단계를 먼저 실행하세요.")
            return

        excel_path  = self.export_excel_path.get().strip()
        delivery_no = self.delivery.get().strip()
        so_no       = self.so.get().strip()
        sheet_name  = self._selected_sheet_name()
        mode_str    = self.mode.get()
        today       = datetime.now().date()

        if not excel_path:
            messagebox.showwarning("주의",
                "수출신고실적 엑셀 파일을 설정에서 선택하세요.")
            return

        rows = []
        for row_id in row_ids:
            vals = self.tree.item(row_id, "values")
            rows.append({
                "destination": sheet_name,
                "purpose":     mode_str,
                "export_date": today,
                "hs_code":     vals[3],
                "part_code":   vals[0],
                "description": vals[1],
                "qty":         int(float(vals[4])),
                "unit_price":  safe_float(vals[5]) or 0.0,
                "so":          so_no,
                "delivery":    delivery_no,
                "origin":      vals[2],
            })

        try:
            n = write_to_export_excel(excel_path, rows)
            self.log(f"✅ 수출신고실적 {n}건 기입 완료 → {os.path.basename(excel_path)}")
            messagebox.showinfo("완료",
                f"수출신고실적에 {n}건 기입 완료!\n"
                f"Delivery: {delivery_no}\n"
                f"파일: {excel_path}")
        except Exception as e:
            self.log(f"[ERROR] 수출신고실적 기입 실패: {e}")
            messagebox.showerror("Error", str(e))


if __name__ == "__main__":
    root = tk.Tk()
    app = CIApp(root)
    root.mainloop()