import os
import re
import json
import shutil
import hashlib
from collections import defaultdict
from datetime import datetime, timedelta

import pdfplumber
from openpyxl import load_workbook

# 영업일 계산용 한국 공휴일: holidays 라이브러리 있으면 자동(설·추석·대체공휴일 포함),
# 없으면 주말만 제외하도록 안전하게 fallback
try:
    import holidays as _holidays
    _yr = datetime.now().year
    _KR_HOLIDAYS = _holidays.SouthKorea(years=range(_yr - 1, _yr + 3))
except Exception:
    _KR_HOLIDAYS = set()


# =========================================================
# ✅ 설정
# =========================================================
PDF_DIR = r"C:\Users\yoongil.chae\OneDrive - Candela\Syneron-Candela Korea - Operation\10. 수입\수입신고필증_자동화"
EXCEL_FILE = r"C:\Users\yoongil.chae\OneDrive - Candela\Syneron-Candela Korea - Operation\10. 수입\수입신고실적(20260211)자동화.xlsx"
SHEET_NAME = r"수입신고실적(20260211)"
BACKUP_DIR = r"C:\Users\yoongil.chae\OneDrive - Candela\Syneron-Candela Korea - Operation\10. 수입\백업"

PROCESSED_DB = os.path.join(os.path.dirname(EXCEL_FILE), "_processed_pdfs.json")

WINDOW_DAYS = 14
SKIP_IF_AMBIGUOUS = False
PROTECTED_COLS = {"본사선적일자", "자재코드", "자재명", "수량"}
# 2026-07-07: "용마입고"(+2영업일 추정치) 컬럼은 "용마 실제 입고"로 개명되어
# 이제 otbi_receipt_updater.py가 Fusion 실데이터로만 채운다. 이 스크립트는
# 더 이상 이 컬럼에 쓰지 않는다(추정치 로직은 검증 기간 동안 잠시 남겨두되
# 결과를 엑셀에 반영하지 않음 — OTBI 파이프라인이 안정화되면 완전히 제거 예정).
PROTECTED_COLS.add("용마입고")


# =========================================================
# ✅ 통관 알림 메일 설정
# =========================================================

# 자재코드(normalize 후) → 이메일에 표시할 제품명
# 9SYS7751-CNDL은 normalize_item_code()가 -CNDL을 제거하므로 키는 9SYS7751
ALERT_CODE_NAMES = {
    "9914-CE-9036": "GMPP",
    "9914-VT-0300": "Vbeam",
    "9SYS7751":     "놀리스",
    "9914-JB-9060": "피코웨이",
}

ALERT_TO = [
    "ckserviceteam@candelamedical.com",
    "CKSalesTeam@candelamedical.com",
    "CK_clinical@syneron.onmicrosoft.com",
]

ALERT_CC = [
    "CK-Operation@syneron.onmicrosoft.com",
]

# 발송 이력 저장 파일 (중복 발송 방지 — 신고번호 기준)
ALERT_SENT_DB = os.path.join(os.path.dirname(EXCEL_FILE), "_alert_sent.json")

# 2026-08-10 사용자 요청: 자동 발송/초안 메일 글꼴을 굴림 -> 맑은 고딕으로 통일.
# 평문(.Body)으로 만들면 Outlook의 평문 기본 글꼴을 따라가므로 HTML로 만든다.
MAIL_FONT_CSS = "font-family:'맑은 고딕',Malgun Gothic,sans-serif;font-size:10.0pt"


def mail_text_to_html(text: str) -> str:
    """평문 본문을 맑은 고딕 HTML로 감싼다.
    HTML은 연속 공백을 하나로 줄이므로, 품목 목록의 정렬이 뭉개지지 않도록
    2칸 이상 연속 공백은 &nbsp;로 보존한다."""
    import html as html_module
    import re as re_module
    esc = html_module.escape(str(text or ""))
    esc = esc.replace("\r\n", "\n").replace("\r", "\n")
    esc = re_module.sub(r"  +", lambda mm: "&nbsp;" * len(mm.group()), esc)
    esc = esc.replace("\t", "&nbsp;&nbsp;&nbsp;&nbsp;")
    return f'<div style="{MAIL_FONT_CSS}">' + esc.replace("\n", "<br>") + "</div>"


# =========================================================
# 유틸
# =========================================================
def log(msg):
    print(f"[{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}] {msg}")

def safe_float(x):
    try:
        if x is None:
            return 0.0
        s = str(x).strip()
        if s == "":
            return 0.0
        return float(s.replace(",", ""))
    except:
        return 0.0

def safe_int(x):
    try:
        if x is None:
            return 0
        s = str(x).strip()
        if s == "":
            return 0
        return int(s.replace(",", ""))
    except:
        return 0

def parse_ymd(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()

    s = str(value).strip()
    if not s:
        return None

    m = re.match(r"(\d{4})[-/](\d{2})[-/](\d{2})", s)
    if m:
        try:
            return datetime(int(m.group(1)), int(m.group(2)), int(m.group(3))).date()
        except:
            return None
    return None

def add_business_days(start_date, n, holidays=_KR_HOLIDAYS):
    """start_date 다음 영업일(주말·공휴일 제외)로 n일 뒤 날짜 반환.
    예: 신고일이 목요일이면 +2영업일 → 다음 주 월요일
    """
    d = start_date
    added = 0
    while added < n:
        d += timedelta(days=1)
        if d.weekday() >= 5:   # 토(5)·일(6)
            continue
        if d in holidays:      # 공휴일
            continue
        added += 1
    return d

def backup_excel():
    os.makedirs(BACKUP_DIR, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = os.path.join(BACKUP_DIR, f"backup_{ts}.xlsx")
    shutil.copy(EXCEL_FILE, backup_path)
    log(f"백업 완료: {backup_path}")

def load_processed_db():
    if os.path.exists(PROCESSED_DB):
        try:
            with open(PROCESSED_DB, "r", encoding="utf-8") as f:
                return json.load(f)
        except:
            return {}
    return {}

def save_processed_db(db):
    with open(PROCESSED_DB, "w", encoding="utf-8") as f:
        json.dump(db, f, ensure_ascii=False, indent=2)

def file_fingerprint(path):
    st = os.stat(path)
    size = st.st_size
    mtime = int(st.st_mtime)

    h = hashlib.sha256()
    with open(path, "rb") as f:
        chunk = f.read(2 * 1024 * 1024)
        h.update(chunk)

    return {"size": size, "mtime": mtime, "sha256_2mb": h.hexdigest()}

def normalize_decl_no(decl_no: str) -> str:
    return (decl_no or "").replace("-", "").strip()

def is_valid_decl_no(x: str) -> bool:
    """신고번호처럼 보이는 값만 True.
    예: 44820-26-720326M / 4482026720326M
    """
    s = normalize_decl_no(x or "")
    return bool(re.fullmatch(r"\d{13,14}[A-Z]", s))

def normalize_item_code(code) -> str:
    """자재코드(Item Code) 매칭용 정규화."""
    if code is None:
        return ""
    # 리스트가 들어오는 경우 방어 처리
    if isinstance(code, list):
        code = code[0] if code else ""
    s = str(code).strip().upper()
    if not s:
        return ""
    if s.startswith("'"):
        s = s[1:]
    s = s.replace("–", "-").replace("—", "-").replace("−", "-")
    parts = s.split()
    if not parts:
        return ""
    s = parts[0]
    for suf in ("-CNDL",):
        if s.endswith(suf):
            s = s[: -len(suf)]
    return s

def append_remark(existing, msg):
    existing = "" if existing is None else str(existing).strip()
    msg = str(msg).strip()
    if not existing:
        return msg
    if msg in existing:
        return existing
    return existing + " | " + msg


# =========================================================
# ✅ 통관 알림 메일 — 발송 이력 관리
# =========================================================
def load_alert_sent_db() -> set:
    if os.path.exists(ALERT_SENT_DB):
        try:
            with open(ALERT_SENT_DB, "r", encoding="utf-8") as f:
                return set(json.load(f))
        except:
            return set()
    return set()

def save_alert_sent_db(sent: set) -> None:
    with open(ALERT_SENT_DB, "w", encoding="utf-8") as f:
        json.dump(sorted(sent), f, ensure_ascii=False, indent=2)


# =========================================================
# ✅ 통관 알림 메일 — 발송 함수
# =========================================================
def send_customs_alert(pdf_filename: str, matched_items: list, sent_db: set) -> bool:
    """
    matched_items: [{"자재코드": ..., "제품명": ..., "수량": float, "신고번호": ...,
                     "용마입고": ..., "AWB": ...}, ...]
    - 신고번호 단위로 메일 1통 발송 (감시 코드 여러 개여도 한 통에 묶음)
    - 품목을 AWB(B/L(AWB)번호)별 섹션으로 묶어 표시 → 동일 품목이 여러 건
      들어와도 AWB로 구분 가능 (예: 10대 + 10대를 20대로 오인하지 않음)
    - 중복 방지: sent_db에 신고번호가 있으면 스킵
    - 발송 성공 시 sent_db에 신고번호 추가
    """
    if not matched_items:
        return False

    신고번호 = matched_items[0]["신고번호"]

    # ✅ 중복 방지
    if 신고번호 in sent_db:
        log(f"[SKIP 메일] 이미 발송한 신고번호: {신고번호}")
        return False

    # AWB → 자재코드별 수량 합산 (AWB 등장 순서 / 코드 등장 순서 유지)
    from collections import OrderedDict
    by_awb = OrderedDict()
    for it in matched_items:
        awb = (str(it.get("AWB") or "").strip()) or "(AWB 미상)"
        code = it["자재코드"]
        if awb not in by_awb:
            by_awb[awb] = OrderedDict()
        if code not in by_awb[awb]:
            by_awb[awb][code] = {"제품명": it["제품명"], "수량": 0.0}
        by_awb[awb][code]["제품명"] = it["제품명"]
        by_awb[awb][code]["수량"] += it["수량"]

    # AWB별 섹션 구성
    sections = []
    grouped_codes = []  # 로그용
    for awb, codes in by_awb.items():
        lines = [f"[AWB: {awb}]"]
        for code, info in codes.items():
            grouped_codes.append(code)
            qty_val = info["수량"]
            qty_str = str(int(qty_val)) if qty_val == int(qty_val) else str(qty_val)
            lines.append(f"  • {info['제품명']} ({code}) : {qty_str}대")
        sections.append("\n".join(lines))

    # 섹션 사이를 구분선으로 연결 (AWB가 2개 이상이면 자연스럽게 분리)
    divider = "─────────────────────"
    item_block = f"\n{divider}\n".join(sections)

    try:
        import win32com.client as _win32
        outlook = _win32.Dispatch("Outlook.Application")
        mail = outlook.CreateItem(0)  # olMailItem

        mail.To = "; ".join(ALERT_TO)
        mail.CC = "; ".join(ALERT_CC)
        용마입고 = matched_items[0].get("용마입고", "")

        mail.Subject = "[통관 완료] 입고 안내"
        mail.HTMLBody = mail_text_to_html(
            f"Dear all,\n\n"
            f"아래 제품이 통관 완료되었습니다.\n"
            f"용마에 {용마입고} 입고되니 업무에 참고 부탁드립니다.\n\n"
            f"{divider}\n"
            f"{item_block}\n"
            f"{divider}\n\n"
            f"감사합니다.\n"
            f"채윤길 드림"
        )
        # 자동 발송하지 않고 임시보관함(Drafts)에 초안으로 저장한다 — 검토 후 직접 [보내기].
        mail.Save()

        sent_db.add(신고번호)  # ✅ 초안 생성 이력 등록(중복 생성 방지)
        log(f"✉️ 알림 메일 초안 저장(검토 후 직접 발송): {grouped_codes} | "
            f"AWB {list(by_awb.keys())} | 신고번호 {신고번호}")
        return True

    except Exception as e:
        log(f"⚠️ 알림 메일 초안 저장 실패: {e}")
        return False


# =========================================================
# ✅ 헤더(컬럼명) 행 자동 탐지
# =========================================================
def find_header_row(ws, required_headers, max_scan_rows=30):
    best_row = None
    best_score = -1

    req = [str(x).strip() for x in required_headers]

    for r in range(1, min(max_scan_rows, ws.max_row) + 1):
        values = []
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if v is None:
                continue
            values.append(str(v).strip())

        if not values:
            continue

        score = sum(1 for h in req if h in values)
        if score > best_score:
            best_score = score
            best_row = r

    if best_row is None or best_score < 4:
        raise ValueError(
            f"헤더 행을 찾지 못했습니다. (스캔 1~{max_scan_rows}행, 매칭최대={best_score})\n"
            f"엑셀 상단에 컬럼명이 있는 행을 확인해 주세요."
        )

    return best_row

def build_header_map(ws, header_row):
    col_map = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(header_row, c).value
        if v is None:
            continue
        name = str(v).strip()
        if name:
            col_map[name] = c
    return col_map


# =========================================================
# PDF 파싱
# =========================================================
def extract_text_all_pages(pdf_path):
    text = ""
    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            t = page.extract_text()
            if t:
                text += t + "\n"
    return text

def _is_plausible_bl(s: str) -> bool:
    if not s:
        return False
    s = s.strip().upper()
    if not re.fullmatch(r"[0-9A-Z\-]+", s):
        return False
    if len(s) >= 6:
        return True
    if "-" in s and len(s) >= 5:
        return True
    return False

def _split_bl_suffix(token: str):
    """B/L 토큰에서 (분할:A) 같은 접미사를 분리. 반환: (base_upper, suffix)
    예: '7386645372(분할:A)' → ('7386645372', '(분할:A)')
    """
    token = token.strip()
    m = re.match(r"([0-9A-Za-z\-]+)(\(.*\))?$", token)
    if m:
        return m.group(1).upper(), (m.group(2) or "")
    return token.upper(), ""

def extract_bl_number(text: str) -> str:
    m = re.search(r"4\s*B/L\(AWB\)번호[^\n]*\n\s*(\S+)", text)
    if m:
        base, suffix = _split_bl_suffix(m.group(1))
        if _is_plausible_bl(base):
            return base + suffix

    m2 = re.search(r"4\s*B/L\(AWB\)번호\s+(\S+)", text)
    if m2:
        base, suffix = _split_bl_suffix(m2.group(1))
        if base != "5" and _is_plausible_bl(base):
            return base + suffix

    m3 = re.search(r"27\s*MASTER\s*B/L번호\s*([0-9A-Z\-]+)", text)
    if m3:
        cand = m3.group(1).strip().upper()
        if _is_plausible_bl(cand):
            return cand

    return ""

def parse_pdf(pdf_path):
    text = extract_text_all_pages(pdf_path)

    decl_m = re.search(r"\b(\d{5}-\d{2}-\d{6,7}[A-Z])\b", text)
    신고번호_raw = decl_m.group(1).strip() if decl_m else ""
    신고번호 = normalize_decl_no(신고번호_raw)

    decl_date_m2 = re.search(r"2\s*신고일.*?(\d{4}/\d{2}/\d{2})", text, re.DOTALL)
    decl_date_m = re.search(r"\b(\d{4}/\d{2}/\d{2})\b", text)
    신고일자 = decl_date_m2.group(1).strip() if decl_date_m2 else (decl_date_m.group(1).strip() if decl_date_m else "")

    신고일_dt = None
    if 신고일자:
        try:
            신고일_dt = datetime.strptime(신고일자, "%Y/%m/%d").date()
        except:
            신고일_dt = None

    arr_m = re.search(r"6\s*입항일\s*([0-9]{4}/[0-9]{2}/[0-9]{2})", text)
    입항일 = arr_m.group(1).strip() if arr_m else ""
    입항일_dt = None
    if 입항일:
        try:
            입항일_dt = datetime.strptime(입항일, "%Y/%m/%d").date()
        except:
            입항일_dt = None

    용마입고 = ""
    if 신고일_dt:
        용마입고 = add_business_days(신고일_dt, 2).strftime("%Y/%m/%d")

    bl_final = extract_bl_number(text)

    hs_m = re.search(r"38\s*세번부호\s*([0-9\.\-]+)", text)
    hs_code = hs_m.group(1).strip() if hs_m else ""

    trade_m = re.search(r"14\s*무역거래처\s+(.+)", text)
    trade_partner = trade_m.group(1).strip() if trade_m else ""

    export_m = re.search(r"25\s*적출국\s+([A-Z]{2})", text)
    export_country = export_m.group(1).strip() if export_m else ""

    duty_rate_m = re.search(r"\n관\s+([0-9]+\.[0-9]+)", text)
    duty_rate = duty_rate_m.group(1).strip() if duty_rate_m else ""

    origin_m = re.search(r"46\s*원산지\s+([A-Z]{2})", text)
    origin_code = origin_m.group(1).strip() if origin_m else ""

    # ============================================================
    # 🌟 추가 1: 세관기재란에서 Delivery Number(복수 가능) 추출
    # - 예: DELIVERY NO.9865516,9863830  (2개/3개/그 이상 가능)
    # - 콤마/공백/줄바꿈 섞여도 숫자 토큰을 모두 추출
    # ============================================================
    del_m = re.search(r"DELIVERY\s*NO\.?\s*([0-9,\s]+)", text, re.IGNORECASE)
    raw_dn = del_m.group(1) if del_m else ""
    delivery_numbers = re.findall(r"\d+", raw_dn) if raw_dn else []
    # 하위호환: 기존 단일 필드도 유지(첫 번째 값)
    pdf_delivery = delivery_numbers[0] if delivery_numbers else ""

    # ============================================================
    # 🌟 수입요건확인(48번 항목) 추출 로직
    #
    # PDF 텍스트 실제 구조:
    #   72-2305012600316303          ← 발급서류번호 (수입요건확인 바로 윗 줄)
    #   수입요건확인 표준통관예정보고서(의료기기)
    #   48
    #   (발급서류명)
    #
    # → 발급서류번호(윗 줄)와 발급서류명(같은 줄)을 모두 추출
    # ============================================================
    requirement = "N"

    # ① 수입요건확인 바로 윗 줄에서 발급서류번호 추출 (숫자+하이픈만 있는 줄)
    doc_num = ""
    m_doc_num = re.search(r"([\d][\d\-]+)\n수입요건확인", text)
    if m_doc_num:
        cand = m_doc_num.group(1).strip()
        if re.match(r"^\d[\d\-]+$", cand) and "-" in cand and len(cand) >= 8:
            doc_num = cand

    # ② 수입요건확인 ~ (발급서류명) 사이에서 발급서류명 추출
    doc_name = ""
    m_req = re.search(r"수입요건확인(.*?)\(발급서류명\)", text, re.DOTALL)
    if m_req:
        raw_val = m_req.group(1)
        val = re.sub(r"\s+", " ", raw_val).strip()
        val = re.sub(r"\b48\b|\b49\b", "", val).strip()
        if val:
            doc_name = val

    # ③ 두 값 조합
    parts = [p for p in [doc_num, doc_name] if p]
    if parts:
        requirement = "\n".join(parts)
    # ============================================================

    # ---------- Item parsing (robust; pdf text order may vary) ----------
    items = []
    
    block_pat = re.compile(r"\(NO\.\s*\d+\)\s*", re.IGNORECASE)
    code_pat = re.compile(r"\b\d{4}-\d{2}-\d{4}\b")
    # FIN101204, SBA101506, 9APP7748-CNDL 같은 변형 코드 대응
    line_code_pat = re.compile(r"(?m)^[ \t]*([0-9A-Z][0-9A-Z\-]{3,25})\s*=")
    numline_pat = re.compile(r"(?P<qty>\d+(?:\.\d+)?)\s+(?P<uom>[A-Z]{1,5})\s+(?P<unit_price>[\d,]+)\s+(?P<amount>[\d,]+)")

    # 1. '란(Section)' 단위로 텍스트 묶기 (여러 페이지에 걸친 동일 '란' 병합)
    lan_dict = {}
    lan_parts = re.split(r"란번호\s*/\s*총란수\s*:\s*(\d{3})\s*/\s*\d{3}", text)
    if len(lan_parts) > 1:
        for i in range(1, len(lan_parts), 2):
            lan_num = lan_parts[i]
            lan_text = lan_parts[i+1]
            if lan_num not in lan_dict:
                lan_dict[lan_num] = ""
            lan_dict[lan_num] += lan_text
    else:
        lan_dict["001"] = text

    for lan_num, lan_text in lan_dict.items():
        # 2. 란 단위의 원산지 추출
        lan_origin = ""
        m_lan_ori = re.search(r"원산지[\s\S]{0,50}?([A-Z]{2})-[A-Z]", lan_text)
        if m_lan_ori:
            lan_origin = m_lan_ori.group(1)
        else:
            m_lan_ori2 = re.search(r"원산지\s*([A-Z]{2})\b", lan_text)
            if m_lan_ori2:
                lan_origin = m_lan_ori2.group(1)

        # 🌟 란 단위 수입요건확인(48번) 추출
        # PDF 구조: 발급서류번호(윗 줄) + 수입요건확인 발급서류명(같은 줄) + (발급서류명)
        lan_requirement = "N"

        # ① 발급서류번호: 수입요건확인 바로 윗 줄 (숫자+하이픈으로만 구성)
        lan_doc_num = ""
        m_doc_num_lan = re.search(r"([\d][\d\-]+)\n수입요건확인", lan_text)
        if m_doc_num_lan:
            cand = m_doc_num_lan.group(1).strip()
            # 발급서류번호는 반드시 하이픈 포함 + 8자 이상 (예: 72-2305012600319107)
            # "00", "48" 같은 짧은 숫자(특수세액 0.00의 끝자리, 필드번호 등)는 제외
            if re.match(r"^\d[\d\-]+$", cand) and "-" in cand and len(cand) >= 8:
                lan_doc_num = cand

        # ② 발급서류명: 수입요건확인 ~ (발급서류명) 사이
        lan_doc_name = ""
        m_req_lan = re.search(r"수입요건확인(.*?)\(발급서류명\)", lan_text, re.DOTALL)
        if m_req_lan:
            raw_val = m_req_lan.group(1)
            lines = [l.strip() for l in raw_val.splitlines() if l.strip()]
            lines = [re.sub(r"^\d+$", "", l).strip() for l in lines]
            lines = [l for l in lines if l]
            if lines:
                lan_doc_name = "\n".join(lines)

        # ③ 두 값 조합
        lan_parts = [p for p in [lan_doc_num, lan_doc_name] if p]
        if lan_parts:
            lan_requirement = "\n".join(lan_parts)

        # 🌟 란 단위 세번부호(HS code) 추출 - 란마다 세번이 다를 수 있음
        # PDF 구조: "38 세번부호 9018.90-9000" 또는 "세번부호\n9018.90-9000"
        lan_hs_code = ""
        m_lan_hs = re.search(r"38\s*세번부호\s*([0-9]{4}\.[0-9]{2}-[0-9]{4})", lan_text)
        if m_lan_hs:
            lan_hs_code = m_lan_hs.group(1).strip()

        # 🌟 란 단위 관세율종가 추출 - 란마다 관세율이 다를 수 있음
        # PDF 구조: "관  6.50 (C가가)" / "관  0.00 (C1가가)" / "관  8.00 (A기가)"
        lan_duty_rate = ""
        m_lan_duty = re.search(r"(?m)^관\s+([0-9]+\.[0-9]+)", lan_text)
        if not m_lan_duty:
            m_lan_duty = re.search(r"\n관\s+([0-9]+\.[0-9]+)", lan_text)
        if m_lan_duty:
            lan_duty_rate = m_lan_duty.group(1).strip()

        # 3. 란 텍스트 내에서 아이템 단위로 파싱
        blocks = block_pat.split(lan_text)
        if len(blocks) > 1:
            blocks = blocks[1:]
        else:
            continue
            
        for blk in blocks:
            blk = blk.strip()
            if not blk:
                continue
            
            mcode = (line_code_pat.search(blk) or code_pat.search(blk))
            if not mcode:
                continue
            code = (mcode.group(1) if mcode.lastindex else mcode.group(0)).strip()
            
            desc = ""
            for line in blk.splitlines():
                if code in line and '=' in line:
                    desc = line.split('=',1)[1].strip()
                    break
                    
            qty = None
            uom = None
            unit_price = None
            amount = None
            for line in blk.splitlines():
                mnum = numline_pat.search(line.replace('\t',' '))
                if mnum:
                    qty = mnum.group('qty')
                    uom = mnum.group('uom')
                    unit_price = mnum.group('unit_price')
                    amount = mnum.group('amount')
                    break
                    
            if qty is None:
                tokens = [t for t in re.split(r"\s+", blk) if t]
                for j,tok in enumerate(tokens):
                    if re.fullmatch(r"[A-Z]{1,5}", tok):
                        cand_nums = []
                        for k in range(max(0,j-6), min(len(tokens), j+7)):
                            if re.fullmatch(r"\d+(?:\.\d+)?", tokens[k]) or re.fullmatch(r"[\d,]+", tokens[k]):
                                cand_nums.append(tokens[k])
                        pure = [n for n in cand_nums if re.fullmatch(r"\d+(?:\.\d+)?", n)]
                        money = [n for n in cand_nums if re.fullmatch(r"[\d,]+", n)]
                        if pure:
                            qty = pure[0]
                            uom = tok
                            if len(money) >= 2:
                                unit_price = money[0]
                                amount = money[1]
                            break
                            
            origin_item = None
            m_ori = re.search(r"46\s*원산지\s*([A-Z]{2})\b", blk)
            if not m_ori:
                m_ori = re.search(r"원산지\s*([A-Z]{2})\b", blk)
            if m_ori:
                origin_item = m_ori.group(1)
                
            items.append({
                '자재코드': code,
                '자재명': desc,
                '수량': qty,
                '수량단위': uom,
                '단가': unit_price,
                '금액': amount,
                '원산지코드': (origin_item or lan_origin or origin_code or ''),
                '요건대상여부': lan_requirement,  # 🌟 란별 수입요건확인 - 품목마다 다를 수 있음
                'Hs code': (lan_hs_code or hs_code),       # 🌟 란별 세번부호 - 란마다 다를 수 있음
                '관세율종가': (lan_duty_rate or duty_rate), # 🌟 란별 관세율 - 란마다 다를 수 있음
            })

    return {
        "신고번호": 신고번호,
        "Delivery Numbers": delivery_numbers,
        "Delivery Number": pdf_delivery,  # 🌟 추가 2: (하위호환) 첫 Delivery Number 저장
        "신고일자": 신고일자,
        "입항일": 입항일,
        "용마입고": 용마입고,
        "B/L번호": bl_final,
        "Hs code": hs_code,
        "무역거래처상호": trade_partner,
        "적출국코드": export_country,
        "관세율종가": duty_rate,
        "원산지코드": origin_code,
        "요건대상여부": requirement,
        "items": items,
        "_신고일_dt": 신고일_dt,
        "_입항일_dt": 입항일_dt
    }


# =========================================================
# 엑셀 업데이트 (서식/필터 유지)
# =========================================================
def cell_value(ws, r, col_map, name):
    return ws.cell(r, col_map[name]).value

def set_cell(ws, r, col_map, name, value):
    ws.cell(r, col_map[name]).value = value


def build_existing_decl_set(ws, header_row, col_map):
    """엑셀에 이미 존재하는 신고번호(정규화)를 한 번에 set으로 구축 (O(N) 1회).
    - ws.cell 루프 대신 iter_rows(values_only=True)로 빠르게 읽음
    """
    decl_col = col_map["신고번호"]
    s = set()
    # values_only=True: 셀 객체가 아니라 값만 받아서 훨씬 빠름
    for (v,) in ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row,
                             min_col=decl_col, max_col=decl_col, values_only=True):
        norm = normalize_decl_no(str(v or ""))
        if norm and is_valid_decl_no(norm):
            s.add(norm)
    return s

def build_row_cache(ws, start_row, end_row, col_map, keys):
    """지정한 컬럼(keys) 값들을 한 번에 읽어 캐싱.
    반환: {key: [row별 값]}  (리스트 인덱스 0 == start_row)
    - openpyxl ws.cell 반복 호출을 크게 줄여 대용량에서 속도 개선
    """
    if end_row < start_row:
        return {k: [] for k in keys}

    cols = [col_map[k] for k in keys]
    minc, maxc = min(cols), max(cols)
    offsets = {k: col_map[k] - minc for k in keys}

    cache = {k: [] for k in keys}
    for row_vals in ws.iter_rows(min_row=start_row, max_row=end_row,
                                 min_col=minc, max_col=maxc, values_only=True):
        for k in keys:
            cache[k].append(row_vals[offsets[k]])
    return cache


def normalize_delivery_no(x) -> str:
    """Delivery Number 정규화.
    - 엑셀에서 9865290.0 형태로 저장되는 값을 9865290으로 통일
    - 공백/콤마 제거
    """
    if x is None:
        return ""
    s = str(x).strip().replace(",", "")
    if not s:
        return ""
    if re.fullmatch(r"\d+(?:\.0+)?", s):
        try:
            return str(int(float(s)))
        except:
            return s
    return s


def normalize_delivery_tokens(x):
    """엑셀/텍스트에 Delivery No가 '9865516,9863830' 같이 여러개 들어올 수 있어
    숫자 토큰들을 리스트로 반환한다. (매칭 키는 토큰 단위로만 사용)
    """
    if x is None:
        return []
    s = str(x).strip()
    if not s:
        return []
    # 콤마/공백/줄바꿈 섞여도 숫자만 추출
    nums = re.findall(r"\d+", s.replace(",", " "))
    out = []
    for n in nums:
        nn = normalize_delivery_no(n)
        if nn:
            out.append(nn)
    # 중복 제거(순서 유지)
    seen = set()
    uniq = []
    for n in out:
        if n not in seen:
            seen.add(n)
            uniq.append(n)
    return uniq

def is_valid_decl_no(x) -> bool:
    """신고번호가 '실제 신고번호' 형태인지 검사.
    예: 44820-26-720326M
    '2' 같은 찌꺼기 값은 False 처리하여 빈칸 채움 대상에서 제외되지 않도록 함.
    """
    if x is None:
        return False
    s = str(x).strip()
    if not s:
        return False
    # 하이픈 제거한 뒤에도 최소 길이/형태 체크
    s_norm = normalize_decl_no(s)
    # 대표 패턴: 5자리-2자리-영숫자(보통 6~8자리 + 문자 0~2)
    return bool(re.fullmatch(r"\d{5}\d{2}[A-Z0-9]{5,}", s_norm, flags=re.I))

def normalize_bl_no(x) -> str:
    if x is None:
        return ""
    return str(x).strip().upper().replace(" ", "")


def log_ambiguous_candidates(row_num, ship_dt, code, candidates):
    ship_s = ship_dt.isoformat() if ship_dt else "None"
    log(f"[AMBIGUOUS → SKIP] row={row_num} 자재코드={code} 본사선적일자={ship_s} 후보={len(candidates)}개")
    for cand in candidates:
        base_dt = cand.get("_입항일_dt") or cand.get("_신고일_dt")
        diff = abs((base_dt - ship_dt).days) if (ship_dt and base_dt) else None
        log(
            f"  - 신고번호={cand.get('신고번호','')} 신고일={cand.get('신고일자','')} "
            f"입항일={cand.get('입항일','')} BL={cand.get('B/L번호','')} diff_days={diff}"
        )

# 🌟 추가 3: Delivery Number 비교 로직 추가 (기존 날짜 매칭은 2순위로 유지)
def choose_best_candidate_for_row(row_ship_dt, excel_delivery, candidates):
    # 1순위: Delivery Number 완벽 매칭
    if excel_delivery:
        matched = []
        ex = normalize_delivery_no(excel_delivery)
        for c in candidates:
            nums = c.get('Delivery Numbers') or [c.get('Delivery Number')]
            nums_norm = [normalize_delivery_no(n) for n in nums if n]
            if ex and ex in nums_norm:
                matched.append(c)
        if matched:
            return matched[0], matched

    # 2순위: 기존 선적일자 기반 매칭 (안전망)
    usable = []
    for cand in candidates:
        base_dt = cand.get("_입항일_dt") or cand.get("_신고일_dt")
        
        if row_ship_dt and base_dt:
            if row_ship_dt > base_dt:
                continue
                
            diff = abs((base_dt - row_ship_dt).days)
            if diff <= WINDOW_DAYS:
                usable.append((diff, cand))
        else:
            usable.append((999999, cand))

    if not usable:
        return None, []

    usable.sort(key=lambda x: x[0])
    best_diff = usable[0][0]
    tied = [c for d, c in usable if d == best_diff]

    if len(tied) == 1:
        return tied[0], tied

    def cand_decl_dt(c):
        dt = c.get("_신고일_dt")
        return dt if dt else datetime.min.date()

    tied.sort(key=cand_decl_dt, reverse=True)
    return tied[0], tied


def build_candidate_index(candidates_by_code):
    """PDF 후보들을 (자재코드 -> Delivery -> [cand]) + (자재코드 -> [cand]) 로 인덱싱.
    ✅ PDF에 Delivery No가 2~3개(콤마 구분) 있어도 각 번호로 모두 인덱싱한다.
    """
    by_code_delivery = {}
    by_code_list = {}
    for code, cands in candidates_by_code.items():
        by_code_list[code] = list(cands)
        dmap = {}
        for c in cands:
            nums = c.get("Delivery Numbers")
            if not nums:
                nums = [c.get("Delivery Number", "")]
            for n in nums:
                d = normalize_delivery_no(n)
                if d:
                    dmap.setdefault(d, []).append(c)
        by_code_delivery[code] = dmap
    return by_code_delivery, by_code_list

def take_candidate(code, ship_dt, excel_delivery, by_code_delivery, by_code_list):
    """
    한 엑셀 행에 대해 최적 후보를 선택.

    ✅ 매칭 키: (Delivery Number + 자재코드) ONLY
    ❌ 날짜 기반 fallback 완전 제거
    """

    code = normalize_item_code(code)
    if not code:
        return None, []

    # 1️⃣ Delivery Number 완전 일치 매칭
    d = normalize_delivery_no(excel_delivery)

    if d:
        cands = by_code_delivery.get(code, {}).get(d, [])
        if cands:
            # Delivery가 같으면 그 안에서만 선택
            best, tied = choose_best_candidate_for_row(ship_dt, d, cands)
            return best, tied

    # 2️⃣ Delivery 불일치 시 절대 안채움
    return None, []

def overwrite_correction_rows(ws, header_row, col_map, decl_no_norm, pdf_data_by_code, now_stamp):
    """수정신고: 같은 신고번호 + 같은 자재코드 행을 업데이트 (기존 행만 수정)"""
    updated = 0
    start_row = header_row + 1
    end_row = ws.max_row

    overwrite_cols = [
        "신고일자", "용마입고", "B/L번호", "Hs code", "요건대상여부",
        "수량단위", "단가", "금액",
        "무역거래처상호", "적출국코드", "관세율종가", "원산지코드"
    ]

    # ✅ 캐싱(읽기용): 신고번호/자재코드/본사선적일자/수량/REMARK
    cache = build_row_cache(
        ws, start_row, end_row, col_map,
        keys=["신고번호", "자재코드", "본사선적일자", "수량", "REMARK"]
    )
    decl_list = cache["신고번호"]
    code_list = cache["자재코드"]
    ship_list = cache["본사선적일자"]
    qty_list = cache["수량"]
    remark_list = cache["REMARK"]

    for i in range(end_row - start_row + 1):
        r = start_row + i

        existing_decl = normalize_decl_no(str(decl_list[i] or ""))
        if existing_decl != decl_no_norm:
            continue

        code = normalize_item_code(code_list[i] or "")
        if not code:
            continue

        item = pdf_data_by_code.get(code)
        if not item:
            continue

        ship_dt = parse_ymd(ship_list[i])
        decl_dt = None
        try:
            if item.get("신고일자"):
                decl_dt = datetime.strptime(item["신고일자"], "%Y/%m/%d").date()
        except:
            decl_dt = None

        # 선적일이 신고일보다 늦으면(이상치) 스킵
        if ship_dt and decl_dt and ship_dt > decl_dt:
            continue

        for col in overwrite_cols:
            if col in PROTECTED_COLS:
                continue

            if col == "금액":
                excel_qty = safe_float(qty_list[i])
                unit_price = safe_float(item.get("단가", 0))
                set_cell(ws, r, col_map, "금액", int(excel_qty * unit_price))
            elif col == "단가":
                set_cell(ws, r, col_map, "단가", item.get("단가", ""))
            else:
                set_cell(ws, r, col_map, col, item.get(col, ""))

        old = remark_list[i]
        new = append_remark(old, f"수정신고 반영({now_stamp})")
        set_cell(ws, r, col_map, "REMARK", new)

        updated += 1

    return updated

def fill_new_rows(ws, header_row, col_map, candidates_by_code_new):
    """엑셀의 '신고번호 빈칸' 행을 채움.

    ✅ 매칭 키: (Delivery Number + 자재코드) ONLY
    ✅ '쪼개진 행' 대응: 같은 Shipment 블록 안에서 (Delivery+자재코드)가 같으면
       PDF 후보를 재사용하여 동일 값으로 채움 (수량/금액 계산에 수량을 쓰는 건 기존대로 유지하되,
       매칭 판단에는 수량을 절대 쓰지 않음)

    Shipment 블록 정의(안전):
    - 아래 앵커 컬럼 중 하나라도 '비어있지 않게' 다시 등장하면서
      (Delivery / 신고번호 / B/L번호 / 신고일자) 값이 이전 컨텍스트와 달라지면 새 블록으로 간주
    - 앵커가 비어있는 연속 행(쪼개진 행)은 이전 블록으로 간주
    """
    updated = 0
    start_row = header_row + 1
    end_row = ws.max_row

    fill_cols = [
        "신고일자", "용마입고", "B/L번호", "신고번호", "Hs code", "요건대상여부",
        "수량단위", "단가", "금액",
        "무역거래처상호", "적출국코드", "관세율종가", "원산지코드"
    ]

    by_code_delivery, by_code_list = build_candidate_index(candidates_by_code_new)

    # ✅ 캐싱(읽기용)
    cache = build_row_cache(
        ws, start_row, end_row, col_map,
        keys=["신고번호", "자재코드", "본사선적일자", "Delivery Number", "수량", "B/L번호", "신고일자"]
    )
    decl_list = cache["신고번호"]
    code_list = cache["자재코드"]
    ship_list = cache["본사선적일자"]
    delivery_list = cache["Delivery Number"]
    qty_list = cache["수량"]
    bl_list = cache["B/L번호"]
    decldate_list = cache["신고일자"]

    # Shipment block state
    block_id = 0
    cur_delivery = ""
    cur_decl = ""
    cur_bl = ""
    cur_decl_date = ""

    # 현재 블록에서 (delivery, code) -> candidate 캐시
    block_candidate_cache = {}

    for i in range(end_row - start_row + 1):
        r = start_row + i

        # --- 현재 행의 앵커 값 추출/정규화
        row_deliveries = normalize_delivery_tokens(delivery_list[i])
        row_delivery = row_deliveries[0] if row_deliveries else ""
        row_decl_raw = decl_list[i]
        row_decl_valid = is_valid_decl_no(row_decl_raw)
        row_decl_norm = normalize_decl_no(str(row_decl_raw).strip()) if row_decl_valid else ""

        row_bl = normalize_bl_no(bl_list[i])
        row_decl_date = (str(decldate_list[i]).strip() if decldate_list[i] is not None else "")
        # (신고일자 셀은 날짜형일 수도 있으니 parse_ymd로 안정화)
        dd = parse_ymd(decldate_list[i])
        if dd:
            row_decl_date = dd.isoformat()

        anchors_present = bool(row_delivery or row_decl_norm or row_bl or row_decl_date)

        # --- 블록 시작 판정
        if anchors_present:
            changed = False
            if row_delivery and row_delivery != cur_delivery:
                changed = True
            if row_decl_norm and row_decl_norm != cur_decl:
                changed = True
            if row_bl and row_bl != cur_bl:
                changed = True
            if row_decl_date and row_decl_date != cur_decl_date:
                changed = True

            if changed:
                block_id += 1
                block_candidate_cache = {}

            # 컨텍스트 업데이트(비어있지 않은 것만)
            if row_delivery:
                cur_delivery = row_delivery
            if row_decl_norm:
                cur_decl = row_decl_norm
            if row_bl:
                cur_bl = row_bl
            if row_decl_date:
                cur_decl_date = row_decl_date

        # --- 채움 대상 여부 (신고번호가 '진짜 신고번호'면 스킵, 아니면 빈칸 취급)
        existing_decl = decl_list[i]
        if is_valid_decl_no(existing_decl):
            continue

        code = normalize_item_code(code_list[i] or "")
        if not code:
            continue

        ship_dt = parse_ymd(ship_list[i])

        # ✅ 매칭 키는 delivery+code ONLY. delivery가 없으면 채울 수 없음.
        delivery_norms = row_deliveries
        if not delivery_norms:
            continue

        # Excel 셀에 Delivery가 여러 개면, 실제로 매칭되는 번호를 찾는다(키는 여전히 1개 Delivery 토큰 단위).
        best = None
        key = None
        for dn in delivery_norms:
            k = (dn, code)
            if k in block_candidate_cache:
                best = block_candidate_cache[k]
                key = k
                break
        if best is None:
            for dn in delivery_norms:
                cand, tied = take_candidate(code, ship_dt, dn, by_code_delivery, by_code_list)
                if cand:
                    best = cand
                    key = (dn, code)
                    block_candidate_cache[key] = best
                    break
        if not best:
            continue

        # --- 쓰기
        for col in fill_cols:
            if col in PROTECTED_COLS:
                continue

            if col == "금액":
                excel_qty = safe_float(qty_list[i])
                unit_price = safe_float(best.get("단가", 0))
                set_cell(ws, r, col_map, "금액", int(excel_qty * unit_price))
            elif col == "단가":
                set_cell(ws, r, col_map, "단가", best.get("단가", ""))
            else:
                set_cell(ws, r, col_map, col, best.get(col, ""))

        updated += 1

    return updated

def update_excel():
    log("자동 업데이트 시작")

    if not os.path.exists(PDF_DIR):
        log(f"❌ PDF 폴더가 존재하지 않습니다: {PDF_DIR}")
        return

    backup_excel()

    wb = load_workbook(EXCEL_FILE)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"시트 '{SHEET_NAME}' 를 찾을 수 없습니다. 현재 시트: {wb.sheetnames}")
    ws = wb[SHEET_NAME]

    # 🌟 추가 6: Delivery Number를 헤더 스캔 목록에 추가
    required_for_header_scan = [
        "본사선적일자", "신고일자", "용마 실제 입고", "B/L번호", "신고번호",
        "자재코드", "자재명", "수량", "REMARK", "Delivery Number"
    ]
    header_row = find_header_row(ws, required_for_header_scan, max_scan_rows=30)
    col_map = build_header_map(ws, header_row)

    log(f"헤더 행: {header_row}")
    log("=== 엑셀 실제 컬럼명 ===")
    log(str(list(col_map.keys())))
    log("========================")

    # 🌟 추가 7: Delivery Number를 필수 컬럼 목록에 추가
    required_cols = [
        "본사선적일자", "신고일자", "용마 실제 입고", "B/L번호", "신고번호", "Hs code", "요건대상여부",
        "자재코드", "자재명", "수량", "수량단위", "단가", "금액",
        "무역거래처상호", "적출국코드", "관세율종가", "원산지코드",
        "REMARK", "Delivery Number"
    ]
    for c in required_cols:
        if c not in col_map:
            raise ValueError(f"엑셀에 '{c}' 컬럼이 없습니다. 현재 컬럼: {list(col_map.keys())}")

    # ✅ 속도 개선: 엑셀의 기존 신고번호를 set으로 한 번만 구축
    existing_decl_set = build_existing_decl_set(ws, header_row, col_map)

    db = load_processed_db()
    pdf_files_all = [f for f in os.listdir(PDF_DIR) if f.lower().endswith(".pdf")]
    pdf_files_all.sort()

    to_process = []
    for f in pdf_files_all:
        p = os.path.join(PDF_DIR, f)
        fp = file_fingerprint(p)
        prev = db.get(f)

        # 파일 자체가 신규이거나 내용이 바뀐 경우
        fp_changed = (
            (prev is None)
            or (prev.get("sha256_2mb") != fp["sha256_2mb"])
            or (prev.get("size") != fp["size"])
        )

        # ✅ 순서 의존 누락 방지:
        # 이전에 '처리됨'으로 기록됐더라도, 그 신고번호가 아직 엑셀에 반영되지
        # 않았으면 재처리한다. (PDF가 채울 빈 행보다 먼저 도착해 '봤지만 못 채운'
        # 경우 — 나중에 빈 행이 생겨도 fingerprint가 같아 영영 누락되던 문제)
        prev_decl = normalize_decl_no(prev.get("decl_no", "")) if prev else ""
        decl_reflected = bool(prev_decl) and (prev_decl in existing_decl_set)

        if fp_changed or (not decl_reflected):
            to_process.append((f, p, fp))

    log(f"PDF 발견: {len(pdf_files_all)}개 / 처리 대상(신규·변경): {len(to_process)}개")
    if not to_process:
        log("신규/변경된 PDF가 없습니다.")
        return

    now_stamp = datetime.now().strftime("%Y-%m-%d %H:%M")

    parsed_by_decl_correction = {}
    candidates_by_code_new = {}

    correction_pdf_count = 0
    new_pdf_count = 0

    # ✅ 알림 메일 발송 이력 로드 (중복 방지)
    alert_sent_db = load_alert_sent_db()

    for f, path, fp in to_process:
        data = parse_pdf(path)
        decl_no = (data.get("신고번호") or "").strip()

        if not decl_no:
            log(f"⚠️ 신고번호 추출 실패: {f} (스킵)")
            continue

        decl_no_norm = normalize_decl_no(decl_no)
        is_correction = (decl_no_norm in existing_decl_set)

        if is_correction:
            correction_pdf_count += 1
            parsed_by_decl_correction.setdefault(decl_no_norm, {})
        else:
            new_pdf_count += 1

        for it in data.get("items", []):
            code = normalize_item_code(it.get("자재코드") or "")
            if not code:
                continue

            candidate = {
                "신고번호": decl_no,
                "Delivery Numbers": data.get("Delivery Numbers", []),
                "Delivery Number": data.get("Delivery Number", ""), # 🌟 추가 8: 추출된 Delivery Number 기억하기
                "신고일자": data.get("신고일자", ""),
                "입항일": data.get("입항일", ""),
                "용마입고": data.get("용마입고", ""),
                "B/L번호": data.get("B/L번호", ""),
                "Hs code": (it.get("Hs code") or data.get("Hs code", "")),
                "요건대상여부": (it.get("요건대상여부") or data.get("요건대상여부", "N")),
                "무역거래처상호": data.get("무역거래처상호", ""),
                "적출국코드": data.get("적출국코드", ""),
                "관세율종가": (it.get("관세율종가") or data.get("관세율종가", "")),
                "원산지코드": (it.get("원산지코드") or data.get("원산지코드", "")),
                "수량단위": it.get("수량단위", ""),
                "단가": it.get("단가", 0),
                "_신고일_dt": data.get("_신고일_dt"),
                "_입항일_dt": data.get("_입항일_dt"),
            }

            # ✅ 후보는 항상 candidates_by_code_new에 적재 (쪼개진 빈 행 채움에 필요)
            candidate["_is_correction"] = bool(is_correction)
            candidates_by_code_new.setdefault(code, []).append(candidate)

            if is_correction:
                # 수정신고: 같은 신고번호 + 같은 자재코드 전부 업데이트 (기존 행 정정용)
                parsed_by_decl_correction[decl_no_norm][code] = candidate

        # ✅ 감시 자재코드 매칭 → 알림 메일 발송
        alert_targets = []
        for it in data.get("items", []):
            code_norm = normalize_item_code(it.get("자재코드") or "")
            if code_norm in ALERT_CODE_NAMES:
                alert_targets.append({
                    "자재코드": code_norm,
                    "제품명":   ALERT_CODE_NAMES[code_norm],
                    "수량":     safe_float(it.get("수량") or 0),
                    "신고번호": decl_no,
                    "용마입고": data.get("용마입고", ""),
                    "AWB":      data.get("B/L번호", ""),
                })
        if alert_targets:
            # 2026-07-07 일시중단: 이 알림 메일이 "+2영업일 추정치"를 인용하는데,
            # 이제 그 추정치를 안 씀. OTBI 실제입고 파이프라인(otbi_receipt_updater.py)이
            # 실제 입고 확인 시점에 알림을 보내도록 옮길 예정 — 그게 검증되면
            # 이 블록과 send_customs_alert()을 완전히 제거할 것.
            pass  # send_customs_alert(f, alert_targets, alert_sent_db)

        fp["decl_no"] = decl_no_norm
        fp["processed_at"] = datetime.now().isoformat(timespec="seconds")
        db[f] = fp

    log(f"[SUMMARY] 정정(수정신고) PDF: {correction_pdf_count}개 / 신규신고 PDF: {new_pdf_count}개")
    if correction_pdf_count > 0:
        log("→ 정정 PDF도 '쪼개진 빈 행' 채움을 위해 후보로는 사용하지만, 날짜 fallback에는 쓰지 않고 Delivery 매칭으로만 채웁니다.")

    correction_updated = 0
    correction_decl_count = 0
    for decl_no, code_map in parsed_by_decl_correction.items():
        u = overwrite_correction_rows(ws, header_row, col_map, decl_no, code_map, now_stamp)
        if u > 0:
            correction_decl_count += 1
            correction_updated += u

    new_updated = 0
    if candidates_by_code_new:
        new_updated = fill_new_rows(ws, header_row, col_map, candidates_by_code_new)

    wb.save(EXCEL_FILE)
    save_processed_db(db)

    # ✅ 알림 메일 발송 이력 저장
    save_alert_sent_db(alert_sent_db)

    log("엑셀 저장 완료")
    log(f"수정신고(같은 신고번호) 정정: 신고 {correction_decl_count}건 / 업데이트 {correction_updated}행")
    log(f"신규(신고번호 빈칸만 채움): 업데이트 {new_updated}행")
    log(f"처리 DB 저장: {PROCESSED_DB}")


if __name__ == "__main__":
    update_excel()
