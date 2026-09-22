"""
Urgent Item Freezer
- K열이 "입고됨"인 행의 I, J, K 수식을 값으로 고정(freeze)
- 수정 전 문제: data_only 캐시가 낡아서 freeze 판단 오류 발생
- 수정 후 해결: Excel COM으로 강제 재계산 후 값 읽기 (COM 불가 시 캐시 fallback)
- 2026-07-23: 입고 알림 메일 발송 기능 제거. OTBI 실제입고 확인 시
              otbi_receipt_updater.py의 send_urgent_arrival_draft가 이미
              같은 알림(초안)을 보내므로 중복 방지 차원에서 삭제함.
"""

import openpyxl
import os
import sys
import subprocess
import tempfile
import shutil
import time
from datetime import datetime

# ── 경로 설정 ──────────────────────────────────────────────────────────────────
EXCEL_FILE_PATH = r'C:\Users\yoongil.chae\OneDrive - Candela\Syneron-Candela Korea - Operation\10. 수입\IR 신청목록 & Instransit(20260303~).xlsx'
root     = r'C:\Users\yoongil.chae\OneDrive - Candela\Syneron-Candela Korea - Operation\10. 수입'
log_path = os.path.join(root, r'작업스케줄러\run_log.txt')

SHEET_NAME    = 'Urgent Item'
COL_PN        = 2   # B
COL_DESC      = 3   # C: 품명(Description)
COL_REQUESTER = 5   # E: 요청자
COL_SHIP_DATE = 9   # I: 해당 물품 최근선적일
COL_YONGMA    = 10  # J: 용마입고날짜
COL_STATUS    = 11  # K: 용마입고여부
FREEZE_VALUE  = '입고됨'

# ── 로그 ───────────────────────────────────────────────────────────────────────
def log(msg):
    line = f'[{datetime.now().strftime("%Y-%m-%d %H:%M:%S")}] {msg}'
    print(line)
    os.makedirs(os.path.dirname(log_path), exist_ok=True)
    with open(log_path, 'a', encoding='utf-8') as f:
        f.write(line + '\n')

def is_formula(value):
    return isinstance(value, str) and value.startswith('=')

# ── 재계산 방법 1: Excel COM (Windows 전용, 가장 정확) ────────────────────────
def recalc_via_excel_com(path):
    """
    Excel COM으로 파일을 열어 강제 재계산 후 저장.
    재계산된 파일을 임시 경로에 복사 후 반환.

    원본 파일은 COM으로 직접 열지 않고 임시 복사본만 연다 — SaveAs 등이
    실패해도 원본 파일에는 잠금이 걸리지 않는다. wb.Close/excel.Quit은
    성공/실패 무관하게 finally에서 항상 호출해 Excel 프로세스가 고아로
    남아 파일을 계속 잠그는 사고(2026-08-05 발생)를 막는다.
    """
    excel = None
    wb = None
    tmp_src = path + '.recalc_src_tmp.xlsx'
    tmp_out = path + '.recalc_tmp.xlsx'
    try:
        import win32com.client
        shutil.copy(path, tmp_src)
        excel = win32com.client.Dispatch('Excel.Application')
        excel.Visible = False
        excel.DisplayAlerts = False
        wb = excel.Workbooks.Open(os.path.abspath(tmp_src))
        excel.CalculateFull()
        wb.SaveAs(os.path.abspath(tmp_out))
        wb.Close(False)
        wb = None
        excel.Quit()
        excel = None
        log('  재계산 방법: Excel COM 성공')
        return tmp_out
    except Exception as e:
        log(f'  Excel COM 실패 ({e}), fallback으로 전환')
        return None
    finally:
        if wb is not None:
            try:
                wb.Close(False)
            except Exception:
                pass
        if excel is not None:
            try:
                excel.Quit()
            except Exception:
                pass
        if os.path.exists(tmp_src):
            try:
                os.remove(tmp_src)
            except Exception:
                pass

# ── 재계산 방법 2: 캐시 그대로 사용 (fallback) ────────────────────────────────
def recalc_via_cache(path):
    """
    data_only=True 캐시를 그대로 사용.
    캐시가 낡을 수 있으므로 경고 출력.
    """
    log('  재계산 방법: data_only 캐시 사용 (Excel COM 없음 — 캐시가 낡으면 일부 행 누락될 수 있음)')
    return None  # None이면 원본 path에서 data_only로 읽음

# ── 메인 freeze 로직 ───────────────────────────────────────────────────────────
def freeze_rows():
    # 1) 재계산된 파일 준비
    recalc_path = recalc_via_excel_com(EXCEL_FILE_PATH)
    use_path    = recalc_path if recalc_path else EXCEL_FILE_PATH

    try:
        # 2) 수식 보존용(원본)과 계산값용(재계산본) 워크북 로드
        wb      = openpyxl.load_workbook(EXCEL_FILE_PATH)           # 수식 그대로
        wb_vals = openpyxl.load_workbook(use_path, data_only=True)  # 계산된 값

        ws      = wb[SHEET_NAME]
        ws_vals = wb_vals[SHEET_NAME]

        # 수식이면 캐시값으로, 캐시도 비었으면 '-' 반환 (로그 표시용 안전 헬퍼)
        def read_display(row_idx, col):
            raw = ws.cell(row=row_idx, column=col).value
            if is_formula(raw):
                cached = ws_vals.cell(row=row_idx, column=col).value
                return cached if cached not in (None, '') else '-'
            return raw if raw not in (None, '') else '-'

        frozen_count = 0

        for row_idx in range(2, ws.max_row + 1):
            k_cell     = ws.cell(row=row_idx, column=COL_STATUS)
            k_raw      = k_cell.value
            k_computed = ws_vals.cell(row=row_idx, column=COL_STATUS).value

            # 입고됨 확정 여부
            is_done = (k_raw == FREEZE_VALUE) or \
                      (is_formula(k_raw) and k_computed == FREEZE_VALUE)

            if not is_done:
                continue

            i_cell  = ws.cell(row=row_idx, column=COL_SHIP_DATE)
            j_cell  = ws.cell(row=row_idx, column=COL_YONGMA)
            changed = False

            if is_formula(i_cell.value):
                i_cell.value = ws_vals.cell(row=row_idx, column=COL_SHIP_DATE).value
                changed = True

            if is_formula(j_cell.value):
                j_cell.value = ws_vals.cell(row=row_idx, column=COL_YONGMA).value
                changed = True

            if is_formula(k_raw):
                k_cell.value = FREEZE_VALUE
                changed = True

            if changed:
                pn = read_display(row_idx, COL_PN)
                log(f'  Row {row_idx} ({pn}): freeze 완료'
                    f'  I={i_cell.value}  J={j_cell.value}')
                frozen_count += 1

        # 3) 저장
        if frozen_count > 0:
            wb.save(EXCEL_FILE_PATH)
            log(f'저장 완료. 총 {frozen_count}개 행 freeze 처리')
        else:
            log('새로 freeze할 행 없음. 파일 변경 없음.')

    finally:
        # 성공/실패 무관하게 임시 파일 반드시 삭제
        if recalc_path and os.path.exists(recalc_path):
            try:
                os.remove(recalc_path)
                log('  임시 재계산 파일 삭제 완료')
            except Exception as e:
                log(f'  임시 파일 삭제 실패 ({e}) — 수동으로 삭제하세요: {recalc_path}')

# ── 실행 ───────────────────────────────────────────────────────────────────────
log('====== Urgent Item Freezer Start ======')

if not os.path.exists(EXCEL_FILE_PATH):
    log(f'FAILED: 파일을 찾을 수 없음 → {EXCEL_FILE_PATH}')
    sys.exit(1)

# 2026-09-08: 통합파일이 그 순간 열려있으면(사용자 Excel/OneDrive 동기화 등)
# PermissionError로 이 단계뿐 아니라 뒤따르는 append/internal_share까지
# 통째로 스킵되던 문제 재발 방지용 재시도.
MAX_RETRIES = 5
RETRY_WAIT_SEC = 60

for attempt in range(1, MAX_RETRIES + 1):
    try:
        freeze_rows()
        log('====== Urgent Item Freezer Done ======')
        break
    except PermissionError:
        if attempt < MAX_RETRIES:
            log(f'파일이 열려 있어 접근 실패 ({attempt}/{MAX_RETRIES}). {RETRY_WAIT_SEC}초 후 재시도.')
            time.sleep(RETRY_WAIT_SEC)
        else:
            log('FAILED: 파일이 Excel에서 열려 있음(재시도 소진). 닫고 다시 실행하세요.')
            sys.exit(1)
    except Exception as e:
        log(f'FAILED: {e}')
        sys.exit(1)

log('')